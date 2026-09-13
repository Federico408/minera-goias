# -*- coding: utf-8 -*-
"""Atualiza o protótipo (v4 -> v5): 03_dim_operacoes e 13_mapas_camadas passam a ser reais
(Base 4) e entra a aba 13b_auditoria_mapas."""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from caminhos import BASE, TMP  # caminhos relativos ao projeto — ver caminhos.py
SRC = f"{BASE}/documentacao/prototipo_bases_consolidadas_v4.xlsx"
OUT = f"{BASE}/documentacao/prototipo_bases_consolidadas_v5.xlsx"
NAVY, TEAL, LBLUE, WHITE, AMBER = "17365D", "1F6D7A", "DCE6F1", "FFFFFF", "C65911"

catalogo = json.load(open(f"{TMP}/_mapas_catalogo.json", encoding="utf-8"))
auditoria = json.load(open(f"{TMP}/_mapas_auditoria.json", encoding="utf-8"))
operacoes = json.load(open(f"{TMP}/_dim_operacoes.json", encoding="utf-8"))
resumo = json.load(open(f"{TMP}/_mapas_resumo.json", encoding="utf-8"))

wb = openpyxl.load_workbook(SRC)


def style_header(ws, title, note, ncol, h=60):
    last = get_column_letter(ncol)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, color=WHITE, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = note
    ws["A2"].font = Font(italic=True, color="404040", size=10)
    ws["A2"].fill = PatternFill("solid", fgColor=LBLUE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = h


def write_table(ws, columns, rows, table_name, header_row=4, wrap=()):
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=c["label"])
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[header_row].height = 30
    for r_i, row in enumerate(rows, start=header_row + 1):
        for c_i, c in enumerate(columns, start=1):
            v = row.get(c["key"])
            v = None if v == "" else v
            cell = ws.cell(row=r_i, column=c_i, value=v)
            if c.get("fmt") and v is not None:
                cell.number_format = c["fmt"]
            if c["key"] in wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    if rows:
        ref = f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}"
        t = Table(displayName=table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c.get("width", 16)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=min(3, len(columns))).coordinate


# --- 03_dim_operacoes -------------------------------------------------------
del wb["03_dim_operacoes"]
ws = wb.create_sheet("03_dim_operacoes")
ws.sheet_view.showGridLines = False
cols = [
    dict(key="operation_id", label="operation_id", width=18), dict(key="processo_anm", label="processo_anm", width=13),
    dict(key="company_id", label="company_id", width=24), dict(key="razao_social", label="razao_social", width=40),
    dict(key="identificacao_empresa", label="identificacao_empresa", width=26),
    dict(key="mineral_ids", label="mineral_ids", width=22), dict(key="mineral_principal", label="mineral_principal", width=22),
    dict(key="municipality_id", label="municipality_id", width=13), dict(key="municipio_nome", label="municipio_nome", width=22),
    dict(key="municipios_intersectados", label="municipios_intersectados", width=30),
    dict(key="qtd_municipios", label="qtd_municipios", width=10, fmt="0"),
    dict(key="fase_atual", label="fase_atual", width=26), dict(key="status_operacional", label="status_operacional", width=38),
    dict(key="categoria_mapa", label="categoria_mapa", width=26),
    dict(key="cfem_2024_2026_brl", label="cfem_2024_2026_brl", width=15, fmt="#,##0.00"),
    dict(key="cfem_2022_2026_brl", label="cfem_2022_2026_brl", width=15, fmt="#,##0.00"),
    dict(key="area_ha_calculada", label="area_ha_calculada", width=13, fmt="#,##0.00"),
    dict(key="area_ha_declarada", label="area_ha_declarada", width=13, fmt="#,##0.00"),
    dict(key="fracao_area_em_go", label="fracao_area_em_go", width=11, fmt="0.0000"),
    dict(key="latitude", label="latitude", width=12, fmt="0.000000"), dict(key="longitude", label="longitude", width=12, fmt="0.000000"),
    dict(key="ultimo_evento", label="ultimo_evento", width=40), dict(key="uso_declarado", label="uso_declarado", width=18),
    dict(key="uf_declarada_sigmine", label="uf_declarada_sigmine", width=14),
    dict(key="qtd_feicoes_originais", label="qtd_feicoes_originais", width=11, fmt="0"),
    dict(key="geometria_reparada", label="geometria_reparada", width=11),
    dict(key="source_ids", label="source_ids", width=40),
]
ordem = {k: i for i, k in enumerate(resumo["rotulos"])}
linhas = sorted(operacoes, key=lambda r: (ordem.get(r["categoria_mapa"], 99), -(r["cfem_2024_2026_brl"] or 0), r["processo_anm"]))
cats = " | ".join(f"{v}: {resumo['categorias'].get(k, 0):,}".replace(",", ".") for k, v in resumo["rotulos"].items())
style_header(ws, "Dimensão — Operações / processos minerários (v5, real — tabela de atributos da camada de mapa CAM_02)",
             f"{len(linhas):,} processos com polígono no SIGMINE que tocam Goiás (fragmentos dissolvidos). ".replace(",", ".")
             + "operation_id = 'OPE_' + número/ano do processo (determinístico). company_id pela ponte do processo (reserva: nome do SIGMINE). "
             + "municipality_id por JUNÇÃO ESPACIAL (município de maior área) — o SIGMINE não tem campo de município; concordância de 99,8% com o texto do Cadastro Mineiro. "
             + "latitude/longitude = representative_point do polígono (sempre dentro dele). Categorias: " + cats,
             len(cols), 90)
write_table(ws, cols, linhas, "DimOperacoesV5")
print("03_dim_operacoes:", len(linhas))

# --- 13_mapas_camadas -------------------------------------------------------
del wb["13_mapas_camadas"]
ws = wb.create_sheet("13_mapas_camadas")
ws.sheet_view.showGridLines = False
cols = [
    dict(key="camada_id", label="camada_id", width=10), dict(key="nome_camada", label="nome_camada", width=30),
    dict(key="descricao", label="descricao", width=46), dict(key="tipo_geometria", label="tipo_geometria", width=18),
    dict(key="n_feicoes", label="n_feicoes", width=10, fmt="#,##0"), dict(key="arquivo_geojson", label="arquivo_geojson (web)", width=44),
    dict(key="tamanho_geojson_mb", label="tamanho_MB", width=10, fmt="0.0"), dict(key="camada_gpkg", label="camada_gpkg (QGIS)", width=50),
    dict(key="crs", label="crs", width=36), dict(key="simplificacao_web", label="simplificacao_web", width=26),
    dict(key="chave_primaria", label="chave_primaria", width=15), dict(key="chaves_de_juncao", label="chaves_de_juncao", width=60),
    dict(key="simbologia", label="simbologia", width=40), dict(key="source_ids", label="source_ids", width=44),
    dict(key="observacao", label="observacao", width=44),
]
style_header(ws, "Base 4 — Catálogo de camadas de mapa (v5, real)",
             "Cada camada é um arquivo GeoJSON (WGS 84, RFC 7946 — direto no Leaflet/Mapbox do Squad 3) e uma camada do GeoPackage "
             "outputs/mapas/minera_goias_mapas_v1.gpkg (SIRGAS 2000, resolução total — para QGIS). Todas carregam as mesmas chaves das Bases 1-3. "
             "Seguindo o guia do Squad 3, ocorrência geológica, processo, operação ativa e projeto futuro têm camadas e símbolos separados; "
             "CAM_05 (ocorrências do RECMIN, aba 06) e CAM_06 (projetos, aba 04) são gravadas por build_mapas_04_06.py, que roda depois dessas abas.",
             len(cols), 66)
write_table(ws, cols, catalogo, "MapasCamadasV5", wrap=("descricao", "chaves_de_juncao", "observacao"))
print("13_mapas_camadas:", len(catalogo))

# --- 13b_auditoria_mapas ----------------------------------------------------
if "13b_auditoria_mapas" in wb.sheetnames:
    del wb["13b_auditoria_mapas"]
ws = wb.create_sheet("13b_auditoria_mapas")
ws.sheet_view.showGridLines = False
cols = [
    dict(key="camada", label="camada", width=28), dict(key="verificacao", label="verificação", width=54),
    dict(key="n", label="n", width=9, fmt="#,##0"), dict(key="resultado", label="resultado", width=70),
    dict(key="decisao", label="decisão", width=70),
]
style_header(ws, "Auditoria — Pente fino das chaves espaciais (Base 4)",
             "Cada linha é uma verificação feita antes de gerar as camadas: sistema de coordenadas, validade e fragmentação dos polígonos, "
             "recorte territorial, junção espacial com municípios (e as discordâncias com o Cadastro Mineiro, uma a uma), atribuição de empresa e "
             "mineral, e — no GeoSGB — junção ponto↔planilha analítica, unidades, limites de detecção e separador decimal.",
             len(cols), 50)
write_table(ws, cols, auditoria, "AuditoriaMapasV5", wrap=("verificacao", "resultado", "decisao"))
print("13b_auditoria_mapas:", len(auditoria))

# --- LEIA-ME -----------------------------------------------------------------
leia = wb["00_LEIA-ME"]
r = leia.max_row + 2
for col in (1, 2):
    leia.cell(row=r, column=col).fill = PatternFill("solid", fgColor=TEAL)
leia.cell(row=r, column=2, value="ATUALIZAÇÃO v5 — Base 4 (mapas) real + 3 correções nas Bases 2 e 3").font = Font(bold=True, size=12, color=WHITE)
r += 1
txt = ("13_mapas_camadas vira o catálogo real das camadas em outputs/mapas/ (GeoJSON para web + GeoPackage), com a auditoria em "
       "13b_auditoria_mapas; 03_dim_operacoes passa a ser real (tabela de atributos da camada de processos). Correções aplicadas ao "
       "desenhar as junções: (1) Base 2 — latitude/longitude do município passam a ser o representative_point (o centro da caixa "
       "envolvente caía fora do próprio município em 20 casos); (2) Base 2 — o PIB do SIDRA vem em MIL reais e estava rotulado como "
       "R$; agora convertido para R$; (3) Base 3 — mineral_id estava como chave interna ('cobre') e não juntava com 01_dim_minerais; "
       "agora MIN_###.")
c = leia.cell(row=r, column=2, value=txt)
c.font = Font(size=10.5)
c.alignment = Alignment(wrap_text=True, vertical="top")
leia.row_dimensions[r].height = 110

wb._sheets.sort(key=lambda s: s.title)
wb.active = 0
wb.save(OUT)
print("SALVO:", OUT)
