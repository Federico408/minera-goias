# -*- coding: utf-8 -*-
"""Abas 06 e 06b (v5 -> v5a_ocorrencias): ocorrências minerais do RECMIN geradas por build_ocorrencias_06.py.
Depois entram a 04 (write_projetos_04), a 12, a 08 e a governança valida tudo."""
import json
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.stdout.reconfigure(encoding="utf-8")
from caminhos import BASE, TMP  # caminhos relativos ao projeto — ver caminhos.py
SRC = f"{BASE}/documentacao/prototipo_bases_consolidadas_v5.xlsx"
OUT = f"{BASE}/documentacao/prototipo_bases_consolidadas_v5a_ocorrencias.xlsx"
HDR = 4
NAVY, TEAL, LBLUE, WHITE = "17365D", "1F6D7A", "DCE6F1", "FFFFFF"

p06 = json.load(open(f"{TMP}/_ocorrencias_06.json", encoding="utf-8"))
p06b = json.load(open(f"{TMP}/_ocorrencias_06b.json", encoding="utf-8"))
res = json.load(open(f"{TMP}/_ocorrencias_06_resumo.json", encoding="utf-8"))
LARG = dict(occurrence_id=12, nome_local=30, mineral_ids=20, mineral_names=30, substancias_original=30, substancias_sem_categoria_anm=16,
            classe_utilitaria=24, importancia=12, status_economico=14, situacao_explotacao=12, motivo_inatividade=14, municipality_id=11,
            municipality_name=20, municipio_informado=20, uf_informada=6, latitude=11, longitude=11, metodo_geoposicionamento=26,
            provincia_mineral=14, rochas_hospedeiras=24, rochas_encaixantes=24, rochas_afloramento=24, morfologia=14, texturas=14,
            tipos_alteracao=18, tipo_afloramento=18, descricao=40, projeto_sgb=30, folha_sgb=18, numero_campo=12, afloramento_id_sgb=12,
            data_cadastro=12, categoria_anm_sobreposta=24, processos_anm_sobrepostos=24, operation_ids_sobrepostos=26, project_ids_sobrepostos=22,
            valor_observado_estimado=11, metodo_estimacao=12, status_validacao=18, responsavel_validacao=30, source_id=36, source_url=60,
            data_acesso=26, periodo_referencia=26, tipo_fonte=9, observacao=60,
            substancia_original=34, substancia_normalizada=34, ocorrencias=11, mineral_id=10, mineral_name=34, tipo_correspondencia=18, justificativa=70)
FMT = dict(latitude="0.000000", longitude="0.000000", afloramento_id_sgb="0", ocorrencias="#,##0")


def br(x, casas=0):
    return f"{x:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def escrever(wb, aba, titulo, nota, dados, tabela, altura):
    pos = wb.sheetnames.index(aba) if aba in wb.sheetnames else None
    if pos is not None:
        del wb[aba]
    ws = wb.create_sheet(aba, pos) if pos is not None else wb.create_sheet(aba)
    ws.sheet_view.showGridLines = False
    cols, linhas = dados["colunas"], dados["linhas"]
    last = get_column_letter(len(cols))
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, color=WHITE, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = nota
    ws["A2"].font = Font(italic=True, color="404040", size=10)
    ws["A2"].fill = PatternFill("solid", fgColor=LBLUE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = altura
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=HDR, column=i, value=c)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[HDR].height = 30
    fmt_col = {cols.index(k): f for k, f in FMT.items() if k in cols}
    for r_i, linha in enumerate(linhas, start=HDR + 1):
        for c_i, v in enumerate(linha):
            if v is None or v == "":
                continue
            cell = ws.cell(row=r_i, column=c_i + 1, value=v)
            if c_i in fmt_col:
                cell.number_format = fmt_col[c_i]
    t = Table(displayName=tabela, ref=f"A{HDR}:{last}{HDR + max(1, len(linhas))}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = LARG.get(c, 16)
    ws.freeze_panes = ws.cell(row=HDR + 1, column=3).coordinate
    print(f"{aba} escrita: {len(linhas)} linhas × {len(cols)} colunas")


wb = openpyxl.load_workbook(SRC)
imp, st, cat = res["importancia"], res["status_economico"], res["categoria_anm"]
fora = ", ".join(f"{uf} {n}" for uf, n in sorted(res["fora_por_uf"].items(), key=lambda x: -x[1]))
nota06 = (
    f"{br(res['dentro_go'])} ocorrências minerais de Goiás do RECMIN (camada 'Ocorrências de Recursos Minerais' do WFS oficial do SGB, baixada em "
    f"{res['data_download'][:10]}; sha256 na 07). A caixa de download trouxe {br(res['baixados'])} pontos; ficaram os que caem dentro de Goiás pela malha do "
    f"IBGE (fora: {fora}). É EVIDÊNCIA GEOLÓGICA, não lista de projetos (guia, seção 9). Importância: depósito {imp.get('Depósito', 0)}, ocorrência "
    f"{imp.get('Ocorrência', 0)}, indício {imp.get('Indício', 0)}, indeterminado {imp.get('Indeterminado', 0)}; status: mina {st.get('Mina', 0)}, garimpo "
    f"{st.get('Garimpo', 0)}, não explotado {st.get('Não explotado', 0)}, indeterminado {st.get('Indeterminado', 0)}. data_cadastro é a data de carga no "
    f"banco, não a de descoberta: 'Mina Ativo(a)' descreve aquele cadastro — só {res['minas_em_operacao_ativa']} das {res['minas_recmin']} 'minas' estão "
    f"dentro de operação ativa na ANM hoje. {br(res['metodo_carta_250k'])} pontos foram posicionados em carta 1:250.000 (erro de centenas de metros). "
    f"mineral_ids lista todas as substâncias da ocorrência (o RECMIN não indica a principal); crosswalk na 06b. Cruzamento com a ANM: "
    f"{br(len([1]) and sum(v for k, v in cat.items() if k != 'fora_de_processo'))} ocorrências dentro de algum processo (pesquisa {cat.get('5_pesquisa', 0)}, "
    f"operação ativa {cat.get('1_operacao_ativa', 0)}), {cat.get('fora_de_processo', 0)} fora; {res['em_projeto_04']} dentro de projetos da 04. "
    "Fora da aba por não terem informação: localizacao_mina ('Mina subterrânea (GPS sem sinal)' em 92%), situacao_garimpo (cópia de situacao_mina), sureg e "
    "origem (valor único), geologo e datum (sempre WGS84).")
escrever(wb, "06_dim_ocorrencias_geologicas", "Dimensão — Ocorrências e depósitos minerais (RECMIN/SGB) de Goiás (v11, real)", nota06, p06, "DimOcorrenciasV11", 120)
ct = res["crosswalk_por_tipo"]
nota06b = (
    f"{res['nomes_substancia']} nomes de substância do RECMIN em Goiás ({br(res['mencoes'])} menções), cada um com a decisão: crosswalk_base1 "
    f"({ct.get('crosswalk_base1', 0)} nomes — o mesmo mapeamento das bases ANM, 01b), sinonimo_recmin ({ct.get('sinonimo_recmin', 0)} — decidido um a um, com "
    f"justificativa) e sem_categoria_anm ({ct.get('sem_categoria_anm', 0)} — sem mineral_id, a ocorrência guarda o nome). Nenhum mineral novo foi criado: os "
    "IDs MIN_### seguem os da 01. Rochas (mármore, quartzito, granito…) recebem a categoria padrão, porque o RECMIN não informa o tipo de uso.")
escrever(wb, "06b_crosswalk_recmin", "Auditoria — Substâncias do RECMIN → mineral_id (v11)", nota06b, p06b, "CrosswalkRecminV11", 60)
wb._sheets.sort(key=lambda s: s.title)
wb.save(OUT)
print("SALVO:", OUT)
