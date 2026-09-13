# -*- coding: utf-8 -*-
"""Atualiza o protótipo (v3 -> v4): substitui 02_dim_empresas e 11_cons_empresa_ano_mineral por
dados REAIS e adiciona a aba de auditoria do cruzamento da chave 'empresa'."""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from caminhos import BASE, TMP  # caminhos relativos ao projeto — ver caminhos.py
SRC = f"{BASE}/documentacao/prototipo_bases_consolidadas_v3.xlsx"
OUT = f"{BASE}/documentacao/prototipo_bases_consolidadas_v4.xlsx"

NAVY, TEAL, LBLUE, WHITE, AMBER = "17365D", "1F6D7A", "DCE6F1", "FFFFFF", "C65911"

emp_dim = json.load(open(f"{TMP}/_emp_dim.json", encoding="utf-8"))
crosswalk = json.load(open(f"{TMP}/_crosswalk_emp.json", encoding="utf-8"))
base3 = json.load(open(f"{TMP}/_base3_rows.json", encoding="utf-8"))

wb = openpyxl.load_workbook(SRC)


def style_header(ws, title, note, ncol, note_height=60):
    last_col = get_column_letter(ncol)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, color=WHITE, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = note
    ws["A2"].font = Font(italic=True, color="404040", size=10)
    ws["A2"].fill = PatternFill("solid", fgColor=LBLUE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = note_height


def write_table(ws, columns, rows, table_name, header_row=4):
    ncol = len(columns)
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=c["label"])
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[header_row].height = 28
    for r_i, row in enumerate(rows, start=header_row + 1):
        for c_i, c in enumerate(columns, start=1):
            val = row.get(c["key"], "")
            if val == "" or val is None:
                val = None
            cell = ws.cell(row=r_i, column=c_i, value=val)
            if c.get("fmt") and val is not None:
                cell.number_format = c["fmt"]
    last_row = header_row + max(1, len(rows))
    if rows:
        tbl = Table(displayName=table_name, ref=f"A{header_row}:{get_column_letter(ncol)}{last_row}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c.get("width", 16)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=min(4, ncol + 1)).coordinate


# ---------------------------------------------------------------------------
# 02_dim_empresas (REAL)
# ---------------------------------------------------------------------------
del wb["02_dim_empresas"]
ws = wb.create_sheet("02_dim_empresas")
ws.sheet_view.showGridLines = False
cols = [
    dict(key="company_id", label="company_id", width=24),
    dict(key="razao_social", label="razao_social", width=46),
    dict(key="cnpj_raiz", label="cnpj_raiz (8 díg.)", width=14),
    dict(key="cnpj_completo_exemplo", label="cnpj_completo_exemplo", width=21),
    dict(key="tipo_pessoa", label="tipo_pessoa", width=13),
    dict(key="identificacao", label="identificacao (como o ID foi atribuído)", width=34),
    dict(key="nomes_alternativos", label="nomes_alternativos", width=42),
    dict(key="qtd_processos", label="qtd_processos", width=12, fmt="0"),
    dict(key="qtd_municipios", label="qtd_municipios", width=12, fmt="0"),
    dict(key="principais_minerais", label="principais_minerais", width=40),
    dict(key="fase_predominante", label="fase_predominante", width=26),
]
style_header(ws, "Dimensão — Empresas e titulares (v4, real)",
             "5.045 titulares com título minerário em Goiás. company_id usa a RAIZ do CNPJ (8 primeiros dígitos), não o CNPJ completo: "
             "o sufixo /0001-19 identifica o estabelecimento (matriz/filial), então agrupar pela raiz é o que faz a Anglo American (42.184.226/0019-69) "
             "ser UMA empresa e não várias. Pessoa física tem CPF mascarado por LGPD ('***370285**') e só pode ser identificada por nome — "
             "esses recebem ID derivado do nome normalizado (sha1 determinístico, estável entre execuções). O CPF que aparece DENTRO do nome "
             "(empresário individual, 'FULANO DE TAL ***456789**') sai mascarado na exibição, sem mudar o ID.",
             len(cols), 84)
write_table(ws, cols, emp_dim, "DimEmpresasV4")
print("02_dim_empresas (real) escrita:", len(emp_dim), "titulares")

# ---------------------------------------------------------------------------
# 02b_crosswalk_empresas (nova)
# ---------------------------------------------------------------------------
if "02b_crosswalk_empresas" in wb.sheetnames:
    del wb["02b_crosswalk_empresas"]
ws2 = wb.create_sheet("02b_crosswalk_empresas")
ws2.sheet_view.showGridLines = False
cols2 = [
    dict(key="fonte", label="fonte", width=24),
    dict(key="valor_original", label="valor_original", width=62),
    dict(key="chave_normalizada", label="chave_normalizada", width=36),
    dict(key="contagem", label="contagem", width=10, fmt="#,##0"),
    dict(key="atribuido", label="atribuído a", width=34),
    dict(key="tipo_correspondencia", label="tipo_correspondência / achado", width=52),
]
style_header(ws2, "Auditoria — Normalização e cruzamento da chave 'empresa' (pente fino)",
             "TRÊS ACHADOS ESTRUTURAIS. (1) O campo CPF_CNPJ do CFEM está IRRECUPERAVELMENTE corrompido: 34.212 das 38.854 linhas vieram em notação "
             "científica ('9,69944E+11') — o CNPJ virou float e perdeu precisão (sobram ~6 dígitos significativos de 14). Ele agrupa de forma estável, "
             "mas NÃO identifica a empresa, e o CFEM não tem coluna de nome. (2) Por isso a ligação CFEM→empresa é feita pela PONTE DO NÚMERO DO PROCESSO "
             "(CFEM.Processo+AnoDoProcesso → Cadastro Mineiro.Processo), que exige canonizar o formato ('860582'+1995 ↔ '860582/1995' ↔ '002393/1935'). "
             "Cobertura: 93,1% das linhas e 90,1% do valor. (3) Identidade pela RAIZ do CNPJ: só 3 raízes têm mais de uma grafia (fundidas) e 5 nomes "
             "apontam para raízes diferentes (NÃO fundidos — são empresas distintas de nome parecido).",
             len(cols2), 90)
write_table(ws2, cols2, crosswalk, "CrosswalkEmpresas")
print("02b_crosswalk_empresas escrita:", len(crosswalk), "linhas")

# ---------------------------------------------------------------------------
# 11_cons_empresa_ano_mineral (REAL)
# ---------------------------------------------------------------------------
del wb["11_cons_empresa_ano_mineral"]
ws3 = wb.create_sheet("11_cons_empresa_ano_mineral")
ws3.sheet_view.showGridLines = False
cols3 = [
    dict(key="company_id", label="company_id", width=24),
    dict(key="razao_social", label="razao_social", width=42),
    dict(key="cnpj_raiz", label="cnpj_raiz", width=12),
    dict(key="tipo_pessoa", label="tipo_pessoa", width=12),
    dict(key="mineral_id", label="mineral_id", width=18),
    dict(key="mineral_nome", label="mineral_nome", width=24),
    dict(key="ano", label="ano", width=8, fmt="0"),
    dict(key="cfem_recolhido_brl", label="cfem_recolhido_brl", width=17, fmt="#,##0.00"),
    dict(key="qtd_processos_titulo", label="qtd_titulos_com_esse_mineral", width=14, fmt="0"),
    dict(key="qtd_processos_total_empresa", label="qtd_processos_total_empresa", width=14, fmt="0"),
    dict(key="qtd_processos_sigmine", label="qtd_processos_sigmine", width=13, fmt="0"),
    dict(key="municipios_atuacao", label="municipios_atuacao", width=44),
    dict(key="qtd_municipios_mineral_ano", label="qtd_municipios_mineral_ano", width=12, fmt="0"),
    dict(key="qtd_municipios", label="qtd_municipios", width=12, fmt="0"),
    dict(key="fase_predominante", label="fase_predominante", width=24),
    dict(key="producao_t", label="producao_t", width=12, fmt="#,##0.000"),
    dict(key="consumo_energia_mwh", label="consumo_energia_mwh", width=15, fmt="#,##0"),
    dict(key="intensidade_mwh_t", label="intensidade_mwh_t", width=14, fmt="#,##0.000"),
    dict(key="identificacao", label="identificacao", width=32),
    dict(key="source_ids", label="source_ids", width=40),
    dict(key="observacao", label="observacao", width=36),
]
# CORRECAO (v5): o build grava a chave interna do mineral ('cobre'); a dimensao 01_dim_minerais usa MIN_###.
# Mesma ordenacao (classe, nome) usada em write_base1_into_workbook.py, para os IDs coincidirem.
_md = json.load(open(f"{TMP}/_min_def.json", encoding="utf-8"))
KEY2ID = {k: f"MIN_{i + 1:03d}" for i, (k, _) in enumerate(sorted(_md.items(), key=lambda kv: (kv[1][1], kv[1][0])))}
for _r in base3:
    _r["mineral_id"] = KEY2ID.get(_r["mineral_id"], _r["mineral_id"])
rows3 = sorted(base3, key=lambda r: (-(r["cfem_recolhido_brl"] or 0), r["razao_social"], r["mineral_nome"], (r["ano"] is None, r["ano"])))
style_header(ws3, "Consolidada — Empresa × Mineral × Ano (Goiás) — v4, dados reais",
             f"{len(rows3)} linhas, 5.045 titulares. A soma de cfem_recolhido_brl BATE EXATAMENTE com o arquivo bruto da CFEM (R$ 867.578.398,91, diferença R$ 0,00): "
             "os 9,9% de CFEM cujo processo não existe no Cadastro Mineiro nem no SIGMINE (série 96xxxx, regimes não cobertos pelos 13 arquivos de título — "
             "inclui R$ 62,4 mi da Mineração Serra Grande) ficam num titular explícito 'COM_NAO_IDENTIFICADO' em vez de sumirem da base. "
             "producao_t / consumo_energia_mwh / intensidade_mwh_t estão VAZIAS de propósito: a produção da ANM (AMB/RAL) só existe agregada por UF×mineral×ano, "
             "nunca por empresa — é o gap conhecido (QC07) que trava a entrega ao Squad 2; a energia depende do pareamento CCEE, que é do Squad 2.",
             len(cols3), 90)
write_table(ws3, cols3, rows3, "ConsEmpresaMineralAnoV4")
print("11_cons_empresa_ano_mineral (real) escrita:", len(rows3), "linhas")

# ---------------------------------------------------------------------------
# LEIA-ME
# ---------------------------------------------------------------------------
leia = wb["00_LEIA-ME"]
r = leia.max_row + 2
b = leia.cell(row=r, column=2, value="ATUALIZAÇÃO v4 — Base 3 (empresas) real")
b.font = Font(bold=True, size=12, color=WHITE)
b.fill = PatternFill("solid", fgColor=TEAL)
leia.cell(row=r, column=1).fill = PatternFill("solid", fgColor=TEAL)
r += 1
txt = ("02_dim_empresas e 11_cons_empresa_ano_mineral agora são reais (5.045 titulares), com a aba de auditoria "
       "02b_crosswalk_empresas. Achado principal: o CPF/CNPJ do CFEM está corrompido (notação científica) e o CFEM "
       "não tem nome de titular — a ligação CFEM→empresa foi feita pela ponte do número do processo, com 93,1% das "
       "linhas e 90,1% do valor casados; o restante fica visível como 'COM_NAO_IDENTIFICADO' para que a soma continue "
       "batendo com o arquivo bruto. Faltam agora as abas 03, 04, 06, 07, 08, 12, 13 e 14 (ainda ilustrativas da v0).")
c = leia.cell(row=r, column=2, value=txt)
c.font = Font(size=10.5)
c.alignment = Alignment(wrap_text=True, vertical="top")
leia.row_dimensions[r].height = 90

wb._sheets.sort(key=lambda s: s.title)
wb.active = 0
wb.save(OUT)
print("SALVO:", OUT)
