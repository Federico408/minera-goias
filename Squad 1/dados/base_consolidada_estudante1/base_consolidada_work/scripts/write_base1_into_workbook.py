# -*- coding: utf-8 -*-
"""Atualiza o protótipo (v0 -> v1): substitui 01_dim_minerais e 09_cons_mineral_ano por dados
REAIS cruzados de ANM Produção Bruta/Beneficiada, CFEM, Cadastro Mineiro e SIGMINE, e adiciona
a aba de auditoria do cruzamento (pente fino). As demais abas (02-08, 10-14) permanecem como
estavam na v0 (ainda ilustrativas, a fazer nas próximas rodadas)."""
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from caminhos import BASE, TMP  # caminhos relativos ao projeto — ver caminhos.py
# modelo da planilha (ENTRADA do pipeline, não apagar). No GitHub fica em documentacao/modelo_demo/: o importador do site
# não carrega pastas "demo", e o v0 só tem dados ilustrativos. Na pasta local ele pode estar direto em documentacao/.
_V0 = [f"{BASE}/documentacao/modelo_demo/prototipo_bases_consolidadas_v0.xlsx", f"{BASE}/documentacao/prototipo_bases_consolidadas_v0.xlsx"]
SRC = next((p for p in _V0 if os.path.exists(p)), _V0[0])
OUT = f"{BASE}/documentacao/prototipo_bases_consolidadas_v1_base1.xlsx"

NAVY, TEAL, LBLUE, WHITE = "17365D", "1F6D7A", "DCE6F1", "FFFFFF"

min_def = json.load(open(f"{TMP}/_min_def.json", encoding="utf-8"))
crosswalk = json.load(open(f"{TMP}/_crosswalk.json", encoding="utf-8"))
base1 = json.load(open(f"{TMP}/_base1_rows.json", encoding="utf-8"))

# --- IDs sequenciais e estaveis: MIN_001.. ordenados por classe, depois nome ---
order = sorted(min_def.items(), key=lambda kv: (kv[1][1], kv[1][0]))
KEY2ID = {key: f"MIN_{i+1:03d}" for i, (key, _) in enumerate(order)}

wb = openpyxl.load_workbook(SRC)


def style_header(ws, title, note, ncol):
    last_col = get_column_letter(ncol)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, color=WHITE, size=14)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = note
    ws["A2"].font = Font(italic=True, color="404040", size=10)
    ws["A2"].fill = PatternFill("solid", fgColor=LBLUE)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 48


def write_table(ws, columns, rows, table_name, header_row=4):
    ncol = len(columns)
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=i, value=c["label"])
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[header_row].height = 28
    for r_i, row in enumerate(rows, start=header_row + 1):
        for c_i, c in enumerate(columns, start=1):
            val = row.get(c["key"], "")
            if val == "" or val is None:
                val = None
            cell = ws.cell(row=r_i, column=c_i, value=val)
            if c.get("fmt") and val is not None:
                cell.number_format = c["fmt"]
    last_row = header_row + max(1, len(rows))
    if rows:
        last_col = get_column_letter(ncol)
        tbl = Table(displayName=table_name, ref=f"A{header_row}:{last_col}{last_row}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c.get("width", 16)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=min(4, ncol + 1)).coordinate


# ---------------------------------------------------------------------------
# 01_dim_minerais (REAL)
# ---------------------------------------------------------------------------
del wb["01_dim_minerais"]
ws = wb.create_sheet("01_dim_minerais")
ws.sheet_view.showGridLines = False
cols = [
    dict(key="mineral_id", label="mineral_id", width=11),
    dict(key="mineral_nome_padronizado", label="mineral_nome_padronizado", width=32),
    dict(key="classe_substancia", label="classe_substancia", width=16),
    dict(key="mineral_grupo", label="mineral_grupo_pai", width=16),
    dict(key="categoria_agregada_anm", label="categoria_agregada_anm", width=14),
    dict(key="status_producao_go", label="status_producao_go", width=20),
    dict(key="observacao", label="observacao", width=90),
]
# CORRECAO (v6): a dimensao mostrava o status DIGITADO na 1a passagem (MIN_DEF), nao o recalculado da
# evidencia -- divergia da aba 09 em 7 minerais (ex.: Potassio aparecia como 'exploracao' tendo CFEM em GO).
# Agora usa o mesmo status calculado da Base 1; mineral sem nenhuma linha GO = sem_titulo_ou_producao_go.
_status_go = {r["mineral_id"]: r["status_producao_go"] for r in base1 if r["uf"] == "GO"}
rows = []
for key, (nome, classe, grupo, agregada, status_prev, obs) in order:
    rows.append(dict(mineral_id=KEY2ID[key], mineral_nome_padronizado=nome, classe_substancia=classe,
                      mineral_grupo=KEY2ID.get(grupo, grupo), categoria_agregada_anm=bool(agregada),
                      status_producao_go=_status_go.get(key, "sem_titulo_ou_producao_go"), observacao=obs))
style_header(ws, "Dimensão — Minerais (v1, real)",
             "67 substâncias, espinha = as 57 categorias oficiais da ANM (Produção Bruta/Beneficiada, todos os estados) + 10 subitens/residuais justificados na aba de auditoria (01b). "
             "IDs sequenciais e estáveis (MIN_001…). status_producao_go é CALCULADO a partir da evidência real: confirmada = há produção ANM ou CFEM em GO; exploracao = só título minerário (Cadastro/SIGMINE/CFEM-titular) em GO; sem_titulo_ou_producao_go = categoria oficial nacional sem nenhuma ocorrência em GO nas bases atuais.",
             len(cols))
write_table(ws, cols, rows, "DimMineraisV1")

print("01_dim_minerais (real) escrita:", len(rows), "minerais")

# ---------------------------------------------------------------------------
# 01b_crosswalk_pente_fino (nova)
# ---------------------------------------------------------------------------
if "01b_crosswalk_pente_fino" in wb.sheetnames:
    del wb["01b_crosswalk_pente_fino"]
ws2 = wb.create_sheet("01b_crosswalk_pente_fino")
ws2.sheet_view.showGridLines = False
cols2 = [
    dict(key="fonte", label="fonte", width=26),
    dict(key="valor_original", label="valor_original (bruto na fonte)", width=30),
    dict(key="valor_normalizado", label="valor_normalizado (sem acento/maiúsculo)", width=32),
    dict(key="contagem", label="contagem", width=10, fmt="#,##0"),
    dict(key="mineral_atribuido", label="mineral_id / nome atribuído", width=30),
    dict(key="tipo_correspondencia", label="tipo_correspondência", width=26),
]
rows2 = []
for r in crosswalk:
    mid = r["mineral_atribuido"]
    key_lookup = None
    for k, (nome, *_r) in min_def.items():
        if nome == mid:
            key_lookup = k
            break
    label = f"{KEY2ID.get(key_lookup,'')} — {mid}" if key_lookup else mid
    rows2.append(dict(fonte=r["fonte"], valor_original=r["valor_original"], valor_normalizado=r["valor_normalizado"],
                       contagem=r["contagem"] if isinstance(r["contagem"], int) else None,
                       mineral_atribuido=label, tipo_correspondencia=r["tipo_correspondencia"]))
rows2.sort(key=lambda r: (r["mineral_atribuido"], r["fonte"]))
style_header(ws2, "Auditoria — Normalização e cruzamento da chave 'mineral' (pente fino)",
             "Toda string bruta de substância/mineral encontrada em ANM Produção Bruta/Beneficiada (categorias oficiais), CFEM, Cadastro Mineiro e SIGMINE (GO), e o mineral_id ao qual foi atribuída. "
             "366 linhas — cobertura de 100% verificada por script (nenhuma ficou sem correspondência). tipo_correspondência: 'identico' = mesmo nome já padronizado; 'fusao_minerio_forma_pura' = 'Minério de X' fundido com 'X' (regra do contrato: não deixar sinônimo virar categoria à parte); "
             "'fusao_mineralogica_ou_grupo' = nome mineralógico (ex.: cassiterita, hematita, barita) reconhecido como o minério/variedade do mineral canônico, ou item agrupado sob uma categoria oficial composta da ANM (ex.: Rochas Ornamentais - Outras, Argilas, Dolomito e Magnesita).",
             len(cols2))
write_table(ws2, cols2, rows2, "CrosswalkPenteFino")

print("01b_crosswalk_pente_fino escrita:", len(rows2), "linhas")

# ---------------------------------------------------------------------------
# 09_cons_mineral_ano (REAL)
# ---------------------------------------------------------------------------
del wb["09_cons_mineral_ano"]
ws3 = wb.create_sheet("09_cons_mineral_ano")
ws3.sheet_view.showGridLines = False
cols3 = [
    dict(key="mineral_id", label="mineral_id", width=11),
    dict(key="mineral_nome", label="mineral_nome", width=26),
    dict(key="classe_substancia", label="classe_substancia", width=15),
    dict(key="uf", label="uf", width=6),
    dict(key="ano", label="ano", width=8, fmt="0"),
    dict(key="producao_rom_t", label="producao_rom_t", width=16, fmt="#,##0.000"),
    dict(key="producao_beneficiada", label="producao_beneficiada", width=17, fmt="#,##0.000"),
    dict(key="unidade_beneficiada", label="unidade_beneficiada", width=12),
    dict(key="valor_venda_beneficiada_brl", label="valor_venda_beneficiada_brl", width=18, fmt="#,##0"),
    dict(key="cfem_qtd_comercializada_t", label="cfem_qtd_comercializada_t", width=17, fmt="#,##0.000"),
    dict(key="cfem_qtd_outras_unidades", label="cfem_qtd_outras_unidades", width=24),
    dict(key="cfem_qtd_linhas_excluidas", label="cfem_qtd_linhas_excluidas", width=12, fmt="0"),
    dict(key="cfem_recolhido_brl", label="cfem_recolhido_brl", width=15, fmt="#,##0"),
    dict(key="qtd_titulares_cfem", label="qtd_titulares_cfem", width=13, fmt="0"),
    dict(key="qtd_processos_cadastro_mineiro", label="qtd_processos_cadastro_mineiro", width=15, fmt="0"),
    dict(key="qtd_titulares_cadastro_mineiro", label="qtd_titulares_cadastro_mineiro", width=15, fmt="0"),
    dict(key="qtd_processos_sigmine", label="qtd_processos_sigmine", width=14, fmt="0"),
    dict(key="categoria_agregada_anm", label="categoria_agregada_anm", width=13),
    dict(key="status_producao_go", label="status_producao_go", width=18),
    dict(key="source_ids", label="source_ids", width=26),
    dict(key="observacao", label="observacao", width=70),
]
rows3 = []
for r in base1:
    r2 = dict(r)
    r2["mineral_id"] = KEY2ID[r["mineral_id"]]
    rows3.append(r2)
rows3.sort(key=lambda r: (r["mineral_id"], r["uf"], (r["ano"] is None, r["ano"])))
style_header(ws3, "Consolidada — Mineral × Ano (Goiás e Brasil) — v1, dados reais",
             f"{len(rows3)} linhas. GO: Produção Bruta/Beneficiada (ANM), CFEM, contagem de processos/titulares do Cadastro Mineiro e de processos do SIGMINE. BR: mesma produção ANM, soma nacional (todas as UF), para comparação. "
             "qtd_processos/titulares (Cadastro/SIGMINE) são uma FOTOGRAFIA atual (não por ano) — repetida em toda linha do mineral; quando o mineral só tem título (sem produção/CFEM), aparece 1 linha com ano em branco. "
             "Nenhum valor foi estimado: célula em branco = fonte não tinha o dado nesse cruzamento (ver 01b para o rastreio completo da chave mineral).",
             len(cols3))
write_table(ws3, cols3, rows3, "ConsMineralAnoV1")

print("09_cons_mineral_ano (real) escrita:", len(rows3), "linhas")

# ---------------------------------------------------------------------------
# 09b_auditoria_cfem_quantidade -- linhas cuja quantidade saiu da soma (valor em R$ mantido)
# ---------------------------------------------------------------------------
_exc = json.load(open(f"{TMP}/_cfem_qtd_excluidas.json", encoding="utf-8"))
for _e in _exc:
    _k = _e.pop("mineral_key")
    _e["mineral_id"], _e["mineral_name"] = KEY2ID.get(_k, _k), min_def[_k][0]
_exc.sort(key=lambda e: -(e["quantidade_convertida"] or 0))
if "09b_auditoria_cfem_quantidade" in wb.sheetnames:
    del wb["09b_auditoria_cfem_quantidade"]
ws4 = wb.create_sheet("09b_auditoria_cfem_quantidade")
ws4.sheet_view.showGridLines = False
cols4 = [
    dict(key="mineral_id", label="mineral_id", width=11), dict(key="mineral_name", label="mineral_name", width=24),
    dict(key="year", label="year", width=7, fmt="0"), dict(key="month", label="month", width=7, fmt="0"),
    dict(key="processo_anm", label="processo_anm", width=13), dict(key="substancia_original", label="substancia_original", width=22),
    dict(key="unidade_original", label="unidade_original", width=10), dict(key="quantidade_bruta", label="quantidade_bruta", width=18),
    dict(key="quantidade_convertida", label="quantidade_convertida", width=22, fmt="#,##0.000"),
    dict(key="unidade_convertida", label="unidade_convertida", width=10),
    dict(key="valor_recolhido_brl", label="valor_recolhido_brl", width=15, fmt="#,##0.00"),
    dict(key="r_por_unidade", label="r_por_unidade", width=14, fmt="0.00000000"),
    dict(key="mediana_r_por_unidade_mineral", label="mediana_r_por_unidade_mineral", width=14, fmt="0.0000"),
    dict(key="motivo", label="motivo", width=70),
]
style_header(ws4, "Auditoria — Quantidade comercializada da CFEM excluída da soma (Base 1)",
             f"{len(_exc)} linhas da CFEM cuja QUANTIDADE ficou fora de cfem_qtd_comercializada_t na aba 09 — o valor em R$ continua somado em cfem_recolhido_brl. "
             "Motivos: (a) quantidade em notação científica, a mesma corrupção de exportação do CPF_CNPJ; (b) quantidade implausível, com R$ por unidade mais de 1.000× abaixo "
             "da mediana do mesmo mineral e unidade (ex.: bilhões de t pagando poucos milhares de reais). Sem essas linhas, a CFEM em t fica na ordem de grandeza da produção do AMB "
             "(Calcário 2022: 14,9 Mt × 15,3 Mt). A linha original continua intacta no arquivo da CFEM.",
             len(cols4))
write_table(ws4, cols4, _exc, "AuditoriaCfemQuantidade")
print("09b_auditoria_cfem_quantidade escrita:", len(_exc), "linhas")

# ---------------------------------------------------------------------------
# 01c_rochas_por_tipo_de_uso -- mapeamento (substância, uso) -> categoria ANM (v7)
# ---------------------------------------------------------------------------
_ru = json.load(open(f"{TMP}/_rochas_uso.json", encoding="utf-8"))
for _e in _ru:
    _k = _e.pop("mineral_key")
    _e["mineral_id"], _e["mineral_name"] = KEY2ID.get(_k, _k), min_def[_k][0]
    if _e["fonte"] != "CFEM (GO)":
        _e["valor_recolhido_brl"], _e["cfem_t"] = None, None
_ru.sort(key=lambda e: (e["substancia"], e["fonte"], -(e["contagem"] or 0)))
if "01c_rochas_por_tipo_de_uso" in wb.sheetnames:
    del wb["01c_rochas_por_tipo_de_uso"]
ws5 = wb.create_sheet("01c_rochas_por_tipo_de_uso")
ws5.sheet_view.showGridLines = False
cols5 = [dict(key="substancia", label="substancia", width=22), dict(key="fonte", label="fonte", width=20),
         dict(key="tipo_de_uso", label="tipo_de_uso", width=22), dict(key="origem_do_uso", label="origem_do_uso", width=40),
         dict(key="mineral_id", label="mineral_id", width=11), dict(key="mineral_name", label="mineral_name", width=30),
         dict(key="regra", label="regra", width=34), dict(key="contagem", label="contagem", width=10, fmt="#,##0"),
         dict(key="valor_recolhido_brl", label="valor_recolhido_brl", width=16, fmt="#,##0.00"),
         dict(key="cfem_t", label="cfem_t", width=16, fmt="#,##0.000")]
style_header(ws5, "Auditoria — Rochas: categoria da ANM definida pelo tipo de uso (v7)",
             "Para rochas, a categoria depende do USO: o mesmo granito é 'Rochas (Britadas) e Cascalho' como brita e 'Rochas Ornamentais' como revestimento. "
             "Uso: Cadastro Mineiro ('Tipo(s) de Uso', alinhado posição a posição com 'Substância(s)' em 100% das linhas de GO) e SIGMINE (campo USO); na CFEM, que não tem uso, "
             "pelo número do processo. Uso não informativo ('Demais substâncias' etc.) → uso mais frequente da rocha nos títulos de GO. Regras: brita/construção civil → Britadas; "
             "revestimento → Ornamentais (granito, mármore, quartzito, gnaisse, sienito, ardósia) ou Ornamentais-Outras; pedra de talhe/decorativa → Ornamentais-Outras; "
             "mármore para cal/cimento/corretivo → Calcário; sienito industrial → Feldspato, Leucita e Nefelina-Sienito; cerâmica vermelha → Argilas; gema → Gemas; demais usos "
             "industriais → Minerais Industriais (Outros); cascalho é sempre Britadas. Validado contra o AMB antes de implementar: Rochas Ornamentais saiu de 35–191× para 0,5–1,3×; "
             "Britadas de 0,26× para 0,6–1,1× (2022–24).",
             len(cols5))
write_table(ws5, cols5, _ru, "RochasPorUso")
print("01c_rochas_por_tipo_de_uso escrita:", len(_ru), "linhas")

# 09c_alertas_cfem_processo -- processo com tonelagem acima do total estadual do AMB (v7)
# ---------------------------------------------------------------------------
_al = json.load(open(f"{TMP}/_cfem_alertas_processo.json", encoding="utf-8"))
for _e in _al:
    _k = _e.pop("mineral_key")
    _e["mineral_id"], _e["mineral_name"] = KEY2ID.get(_k, _k), min_def[_k][0]
if "09c_alertas_cfem_processo" in wb.sheetnames:
    del wb["09c_alertas_cfem_processo"]
ws6 = wb.create_sheet("09c_alertas_cfem_processo")
ws6.sheet_view.showGridLines = False
cols6 = [dict(key="mineral_id", label="mineral_id", width=11), dict(key="mineral_name", label="mineral_name", width=28),
         dict(key="year", label="year", width=7, fmt="0"), dict(key="processo_anm", label="processo_anm", width=13),
         dict(key="substancias", label="substancias", width=24), dict(key="usos", label="usos", width=20),
         dict(key="cfem_t_processo", label="cfem_t_processo", width=16, fmt="#,##0"),
         dict(key="amb_total_uf_t", label="amb_total_uf_t", width=16, fmt="#,##0"),
         dict(key="razao", label="razao", width=9, fmt="#,##0.00"), dict(key="severidade", label="severidade", width=10),
         dict(key="valor_recolhido_brl", label="valor_recolhido_brl", width=15, fmt="#,##0.00"),
         dict(key="r_por_t", label="r_por_t", width=11, fmt="0.0000"),
         dict(key="mediana_r_por_t_mineral", label="mediana_r_por_t_mineral", width=13, fmt="0.0000"),
         dict(key="observacao", label="observacao", width=70)]
_n_alta = sum(1 for _e in _al if _e["severidade"] == "alta")
style_header(ws6, "Alertas — Processo com tonelagem na CFEM acima do total do AMB para o estado (v7)",
             f"{len(_al)} processo-anos em que UM único processo declara na CFEM mais toneladas comercializadas do que o AMB registra para Goiás inteiro na mesma categoria "
             f"(maior entre produção bruta e beneficiada em t). Severidade ALTA ({_n_alta}): mais de 2× o total estadual, ou R$/t mais de 10× abaixo da mediana do mineral "
             f"(indício de tonelagem inflada). MODERADA ({len(_al) - _n_alta}): entre 1 e 2× com R$/t compatível — pode ser venda de estoque, AMB preliminar ou diferença de base. "
             "As linhas NÃO foram excluídas da soma (o AMB também é declaratório): ficam sinalizadas para revisão da declaração ou do tipo de uso.", len(cols6))
write_table(ws6, cols6, _al, "AlertasCfemProcesso")
print("09c_alertas_cfem_processo escrita:", len(_al), "linhas")

# ---------------------------------------------------------------------------
# Nota de atualizacao no LEIA-ME
# ---------------------------------------------------------------------------
leia = wb["00_LEIA-ME"]
r = leia.max_row + 2
banner = leia.cell(row=r, column=2, value="ATUALIZAÇÃO v1 (Base 1 real)")
banner.font = Font(bold=True, size=12, color=WHITE)
banner.fill = PatternFill("solid", fgColor=TEAL)
leia.cell(row=r, column=1).fill = PatternFill("solid", fgColor=TEAL)
r += 1
texto = ("01_dim_minerais e 09_cons_mineral_ano deixaram de ser ilustrativas: agora são o cruzamento REAL de "
         "ANM Produção Bruta/Beneficiada (57 categorias oficiais), CFEM, Cadastro Mineiro e SIGMINE (Goiás). "
         "A normalização da chave 'mineral' foi conferida item a item — ver a nova aba 01b_crosswalk_pente_fino "
         "(366 linhas, cobertura de 100%, com o motivo de cada fusão/ambiguidade). As demais abas (02–08, 10–14) "
         "ainda são a maquete ilustrativa da v0 e serão substituídas nas próximas rodadas.")
cell = leia.cell(row=r, column=2, value=texto)
cell.font = Font(size=10.5)
cell.alignment = Alignment(wrap_text=True, vertical="top")
leia.row_dimensions[r].height = 60

def sort_key(name):
    # 00_LEIA-ME, 01_dim_minerais, 01b_crosswalk..., 02_..., ... mantém ordem numerica/alfabetica esperada
    return name

wb._sheets.sort(key=lambda ws: ws.title)
wb.active = 0
wb.save(OUT)
print("SALVO:", OUT)
