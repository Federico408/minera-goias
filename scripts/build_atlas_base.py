# -*- coding: utf-8 -*-
"""Atualiza o atlas (data/atlas/atlas.json e processes.json) com a base consolidada do Squad 1 / Estudante 1.

Mantém exatamente o formato lido por public/atlas.js. Vêm da base consolidada (v13):
  municipalities   malha IBGE 2025 (camada municipios_go da Base 4), processos do SIGMINE que tocam cada município e CFEM 2022–2026
  cfem             CFEM por município e ano para os 246 municípios (aba 08 da planilha; 2026 até o último mês do arquivo)
  production       quantidade comercializada em t e CFEM por substância × município em 2025 — é CFEM declarada, não produção
  cfem_years       CFEM por ano: valor, linhas do arquivo, municípios e titulares distintos (company_id)
  cfem_comparable  CFEM de janeiro a julho de cada ano
  beneficiated     valor de venda da produção beneficiada de Goiás (Anuário Mineral Brasileiro)
  investment       investimento declarado em pesquisa mineral em Goiás (arquivo da ANM catalogado na aba 07)
  processes.json   os processos do SIGMINE em Goiás (camada processos_minerarios_go), geometria simplificada e quantizada
Continuam do retrato anterior (eliel.html), porque a base consolidada não os cobre: energy, energy_months (CCEE) e dams (SIGBM).

Uso, fora da VPS (precisa de geopandas/pyogrio, shapely >= 2.1 e openpyxl):
    python scripts/build_atlas_base.py --base "<pasta do projeto do Squad 1>"
A pasta do projeto contém outputs/mapas/minera_goias_mapas_v1.gpkg, documentacao/prototipo_bases_consolidadas_v13.xlsx e
dados/ANM/investimento_pesquisa/InvestimentoPesquisaMineralUf.csv — é o que o pipeline de
Squad 1/dados/base_consolidada_estudante1/ gera e lê.
"""
import argparse
import base64
import csv
import hashlib
import io
import json
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import numpy as np
import openpyxl
import pyogrio
import shapely

ROOT = Path(__file__).resolve().parents[1]
ANOS = [2022, 2023, 2024, 2025, 2026]  # os mesmos passos do seletor de ano do atlas (YEARS em public/atlas.js)
GRUPOS = ["Lavra / extração autorizada", "Pesquisa autorizada", "Requerimento em análise", "Disponibilidade", "Outros"]
GRUPO_DA_FASE = {  # a mesma divisão do retrato anterior; public/atlas.js tem uma cor para cada um dos cinco grupos
    "CONCESSÃO DE LAVRA": 0, "LAVRA GARIMPEIRA": 0, "LICENCIAMENTO": 0, "REGISTRO DE EXTRAÇÃO": 0,
    "AUTORIZAÇÃO DE PESQUISA": 1,
    "DIREITO DE REQUERER A LAVRA": 2, "REQUERIMENTO DE LAVRA": 2, "REQUERIMENTO DE LAVRA GARIMPEIRA": 2,
    "REQUERIMENTO DE LICENCIAMENTO": 2, "REQUERIMENTO DE PESQUISA": 2, "REQUERIMENTO DE REGISTRO DE EXTRAÇÃO": 2,
    "APTO PARA DISPONIBILIDADE": 3, "DISPONIBILIDADE": 3,
}
MANTIDAS = ["energy", "energy_months", "dams"]


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for bloco in iter(lambda: f.read(1 << 20), b""):
            h.update(bloco)
    return h.hexdigest()


def b64(array, dtype):
    return base64.b64encode(np.asarray(array, dtype=dtype).tobytes()).decode("ascii")


def ler_planilha(caminho):
    """Uma passada na aba 08 (fato longo): CFEM e venda beneficiada; e a data do SIGMINE e o total da aba 10 para conferência."""
    wb = openpyxl.load_workbook(caminho, read_only=True)
    linhas = wb["08_fato_producao_energia"].iter_rows(min_row=4, values_only=True)
    ix = {c: i for i, c in enumerate(next(linhas))}
    cfem_mun_ano, cfem_ano, jan_jul = defaultdict(int), defaultdict(int), defaultdict(int)  # centavos: somas exatas
    registros, muns_ano, titulares_ano = Counter(), defaultdict(set), defaultdict(set)
    producao = defaultdict(lambda: [0.0, 0, set()])  # (SUBSTÂNCIA, município) em 2025 -> [t, centavos, titulares]
    venda, subs_venda, periodos = defaultdict(float), defaultdict(set), set()
    for r in linhas:
        fonte, metrica = r[ix["source_id"]], r[ix["metrica"]]
        if fonte == "SRC_ANM_CFEM":
            ano, mun, periodo = int(r[ix["year"]]), str(r[ix["municipality_id"]]), str(r[ix["periodo_referencia"]])
            substancia = str(r[ix["mineral_name"]] or "NÃO CLASSIFICADO").upper()
            if metrica == "cfem_recolhido":
                c = round(float(r[ix["valor_tratado"]] or 0) * 100)
                periodos.add(periodo)
                cfem_mun_ano[(mun, ano)] += c
                cfem_ano[ano] += c
                registros[ano] += 1
                muns_ano[ano].add(mun)
                titulares_ano[ano].add(r[ix["company_id"]])
                if int(periodo[5:7]) <= 7:
                    jan_jul[ano] += c
                if ano == 2025:
                    producao[(substancia, mun)][1] += c
                    producao[(substancia, mun)][2].add(r[ix["company_id"]])
            elif (metrica == "quantidade_comercializada_cfem" and ano == 2025 and r[ix["unidade_padrao"]] == "t"
                  and r[ix["status_validacao"]] != "quantidade_excluida_da_soma_09b"):  # a mesma exclusão da aba 09b
                producao[(substancia, mun)][0] += float(r[ix["valor_tratado"]] or 0)
        elif fonte == "SRC_ANM_PROD_BENEF" and metrica == "valor_venda_beneficiada" and r[ix["uf"]] == "GO":
            v, ano = float(r[ix["valor_tratado"]] or 0), int(r[ix["year"]])
            venda[ano] += v
            if v > 0:
                subs_venda[ano].add(r[ix["mineral_id"]])
    fontes = wb["07_dim_fontes"].iter_rows(min_row=4, values_only=True)
    cab07 = list(next(fontes))
    sigmine = next(dict(zip(cab07, r)) for r in fontes if r and r[0] == "SRC_ANM_SIGMINE")
    aba10 = wb["10_cons_municipio_ano"].iter_rows(min_row=4, values_only=True)
    cab10 = list(next(aba10))
    total10 = sum(round(float(r[cab10.index("cfem_recolhido_brl")] or 0) * 100) for r in aba10 if r and r[0])
    wb.close()
    return dict(cfem_mun_ano=cfem_mun_ano, cfem_ano=cfem_ano, jan_jul=jan_jul, registros=registros, muns_ano=muns_ano,
                titulares_ano=titulares_ano, producao=producao, venda=venda, subs_venda=subs_venda, periodos=sorted(periodos),
                data_sigmine=sigmine["data_arquivo_local"], total10=total10)


def investimento(caminho):
    texto = Path(caminho).read_bytes().decode("cp1252")
    linhas = list(csv.reader(io.StringIO(texto), delimiter=";"))
    cab = linhas[0]
    valores = [i for i, c in enumerate(cab) if c.startswith("Valor")]
    numero = lambda s: float(s.strip().replace(".", "").replace(",", ".")) if s.strip() not in ("", "-") else 0.0
    total = defaultdict(float)
    for r in linhas[1:]:
        if r and r[0].strip() == "GO":
            total[int(r[cab.index("Ano")])] += sum(numero(r[i]) for i in valores)
    return [{"Ano": a, "TOTAL": round(total[a], 2)} for a in sorted(total)]


def aneis_municipio(geom, casas=5):
    partes = list(geom.geoms) if geom.geom_type == "MultiPolygon" else [geom]
    aneis = []
    for p in partes:
        for anel in [p.exterior, *p.interiors]:
            pts = []
            for lon, lat in anel.coords:
                ponto = [round(lat, casas), round(lon, casas)]
                if not pts or pts[-1] != ponto:
                    pts.append(ponto)
            if pts[0] != pts[-1]:
                pts.append(pts[0])
            if len(pts) >= 4:
                aneis.append(pts)
    return aneis


def municipios(gpkg, dados, tolerancia):
    mun = pyogrio.read_dataframe(gpkg, layer="municipios_go", columns=["municipality_id", "municipio_nome", "qtd_processos_sigmine"])
    geoms = mun.geometry.values
    try:  # simplifica a cobertura inteira: vizinhos continuam com a mesma fronteira, sem frestas nem sobreposição
        simples = shapely.coverage_simplify(geoms, tolerancia)
    except Exception as e:
        print(f"coverage_simplify indisponível ({e}); usando simplify por polígono", file=sys.stderr)
        simples = shapely.simplify(geoms, tolerancia, preserve_topology=True)
    saida = []
    for (_, m), g in zip(mun.iterrows(), simples):
        code = str(m["municipality_id"])
        centavos = sum(dados["cfem_mun_ano"].get((code, a), 0) for a in ANOS)
        aneis = aneis_municipio(g)
        assert aneis, f"município sem geometria depois da simplificação: {code}"
        saida.append({"code": code, "name": m["municipio_nome"], "rings": aneis,
                      "processes": int(m["qtd_processos_sigmine"] or 0), "cfem_total": centavos / 100, "_centavos": centavos})
    return sorted(saida, key=lambda x: x["code"])


def processos(gpkg, tolerancia):
    pr = pyogrio.read_dataframe(gpkg, layer="processos_minerarios_go",
                                columns=["processo_anm", "fase_atual", "substancias_sigmine", "area_ha_declarada", "area_ha_calculada"])
    originais = shapely.force_2d(pr.geometry.values)
    simples = shapely.simplify(originais, tolerancia, preserve_topology=True)
    ruins = shapely.is_empty(simples) | ~shapely.is_valid(simples)
    simples[ruins] = originais[ruins]
    lon0, lat0, lon1, lat1 = shapely.total_bounds(originais)
    q = lambda lon, lat: (int(round((lon - lon0) / (lon1 - lon0) * 65535)), int(round((lat1 - lat) / (lat1 - lat0) * 65535)))

    def aneis(geom):
        partes = list(geom.geoms) if geom.geom_type == "MultiPolygon" else [geom]
        saida = []
        for p in partes:
            for anel in [p.exterior, *p.interiors]:
                pts = []
                for lon, lat in anel.coords:
                    k = q(lon, lat)
                    if not pts or pts[-1] != k:
                        pts.append(k)
                if len(pts) > 1 and pts[0] == pts[-1]:
                    pts.pop()
                if len(pts) >= 3:
                    saida.append(pts)
        return saida

    xy, inicio, dono = [], [0], []
    for i, (g, orig) in enumerate(zip(simples, originais)):
        lista = aneis(g) or aneis(orig)
        if not lista:  # polígono menor que o passo da quantização (~12 m): mantém o contorno original, mesmo com pontos repetidos
            lista = [[q(lon, lat) for lon, lat in orig.geoms[0].exterior.coords[:-1]] if orig.geom_type == "MultiPolygon"
                     else [q(lon, lat) for lon, lat in orig.exterior.coords[:-1]]]
        for pts in lista:
            for x, y in pts:
                xy.extend((x, y))
            inicio.append(inicio[-1] + len(pts))
            dono.append(i)
    fases = pr["fase_atual"].fillna("").astype(str)
    subs = pr["substancias_sigmine"].fillna("").astype(str)
    d_fase, d_subs = sorted(set(fases)), sorted(set(subs))
    area = pr["area_ha_declarada"].where(pr["area_ha_declarada"].notna(), pr["area_ha_calculada"]).fillna(0.0)
    grupo = [GRUPO_DA_FASE.get(f, 4) for f in fases]
    n = len(pr)
    assert n < 65536 and len(d_fase) < 256 and len(d_subs) < 65536
    resumo = []
    for gi, nome in enumerate(GRUPOS):
        sel = [k for k in range(n) if grupo[k] == gi]
        resumo.append({"grupo": nome, "n": len(sel), "area": round(float(area.iloc[sel].sum()), 2)})
    return {"n": n, "xy": b64(xy, "<u2"), "ringStart": b64(inicio, "<u4"), "ringPoly": b64(dono, "<u2"),
            "g": b64(grupo, "u1"), "fase": b64([d_fase.index(f) for f in fases], "u1"), "subs": b64([d_subs.index(s) for s in subs], "<u2"),
            "area": b64(area.to_numpy(), "<f4"), "processo": "\x01".join(pr["processo_anm"].astype(str)),
            "dFase": d_fase, "dSubs": d_subs, "grupos": GRUPOS, "resumo": resumo,
            "bounds": {"lon0": round(float(lon0), 6), "lon1": round(float(lon1), 6), "lat0": round(float(lat0), 6), "lat1": round(float(lat1), 6)}}


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--base", required=True, help="pasta do projeto do Squad 1 (com outputs/, documentacao/ e dados/)")
    ap.add_argument("--planilha", help="planilha consolidada (padrão: <base>/documentacao/prototipo_bases_consolidadas_v13.xlsx)")
    ap.add_argument("--tolerancia-municipios", type=float, default=0.003, help="graus; padrão 0,003 (~330 m)")
    ap.add_argument("--tolerancia-processos", type=float, default=0.0004, help="graus; padrão 0,0004 (~45 m)")
    args = ap.parse_args()
    base = Path(args.base)
    planilha = Path(args.planilha) if args.planilha else base / "documentacao" / "prototipo_bases_consolidadas_v13.xlsx"
    gpkg = base / "outputs" / "mapas" / "minera_goias_mapas_v1.gpkg"
    invest = base / "dados" / "ANM" / "investimento_pesquisa" / "InvestimentoPesquisaMineralUf.csv"
    destino = ROOT / "data" / "atlas"
    anterior = json.loads((destino / "atlas.json").read_text(encoding="utf-8"))
    retido = anterior["meta"].get("retained_from_artifact") or {"artifact": anterior["meta"]["artifact"], "sha256": anterior["meta"]["sha256"],
                                                                "periods": {k: anterior["meta"]["periods"][k] for k in ("energy", "dams")}}
    retido["layers"] = MANTIDAS

    dados = ler_planilha(planilha)
    anos = sorted(dados["cfem_ano"])
    assert set(anos) <= set(ANOS), f"anos de CFEM fora do seletor do atlas: {anos}"
    mun = municipios(gpkg, dados, args.tolerancia_municipios)
    assert len(mun) == 246
    total_mun = sum(m.pop("_centavos") for m in mun)
    total_ano = sum(dados["cfem_ano"].values())
    assert total_mun == total_ano == dados["total10"], (total_mun, total_ano, dados["total10"])
    for m in mun:
        for anel in m["rings"]:
            for lat, lon in anel:
                assert -19.51 < lat < -12.38 and -53.26 < lon < -45.89, (m["code"], lat, lon)
    por_nome = {m["code"]: m["name"] for m in mun}

    prod, agregado = defaultdict(dict), defaultdict(lambda: [0.0, 0, set()])
    for (sub, code), (t, c, titulares) in dados["producao"].items():
        if c <= 0 and t <= 0:
            continue
        prod[sub][code] = [round(t, 2), round(c / 100, 2), len(titulares)]
        agregado[sub][0] += t
        agregado[sub][1] += c
        agregado[sub][2].add(code)
    subs = sorted(({"sub": s, "ton": round(v[0], 2), "cfem": round(v[1] / 100, 2), "municipios": len(v[2])} for s, v in agregado.items()),
                  key=lambda x: (-x["municipios"], -x["cfem"]))
    data_sigmine = "/".join(reversed(str(dados["data_sigmine"]).split("-")))
    ultimo = dados["periodos"][-1]
    packet = {
        "meta": {
            "artifact": "base consolidada do Squad 1 / Estudante 1 — prototipo_bases_consolidadas_v13.xlsx",
            "sha256": sha256(planilha),
            "integrated_on": date.today().isoformat(),
            "status": "snapshot_unvalidated",
            "note": ("Retrato da base consolidada do Squad 1 (v13), gerado por scripts/build_atlas_base.py; não é consulta em tempo real à ANM. "
                     "Energia (CCEE) e barragens (SIGBM) seguem do retrato anterior (eliel.html). Geometrias simplificadas e quantizadas para "
                     "o mapa: não usar como limite cadastral."),
            "sources": ["ANM / CFEM", "Anuário Mineral Brasileiro", "Cadastro Mineiro", "SIGMINE", "IBGE — malha municipal 2025",
                        "SIGBM (retrato anterior)", "CCEE (retrato anterior)"],
            "periods": {"cfem": f"{dados['periodos'][0]} a {ultimo}", "production": "2025 — quantidade comercializada declarada na CFEM",
                        "energy": retido["periods"]["energy"], "dams": retido["periods"]["dams"],
                        "processes": f"arquivo do SIGMINE de {data_sigmine}"},
            "base": {"planilha": "Squad 1/dados/base_consolidada_estudante1/documentacao/prototipo_bases_consolidadas_v13.xlsx",
                     "sha256_planilha": sha256(planilha), "sha256_gpkg": sha256(gpkg), "sha256_investimento": sha256(invest)},
            "retained_from_artifact": retido,
        },
        "municipalities": mun,
        "cfem": {"anos": ANOS, "linhas": [{"Município": por_nome[m["code"]], **{str(a): dados["cfem_mun_ano"].get((m["code"], a), 0) / 100 for a in ANOS}}
                                          for m in mun]},
        "production": {"subs": subs, "dados": {s["sub"]: prod[s["sub"]] for s in subs}},
        "energy": anterior["energy"],
        "dams": anterior["dams"],
        "cfem_years": [{"Ano": a, "valor": dados["cfem_ano"][a] / 100, "registros": dados["registros"][a],
                        "municipios": len(dados["muns_ano"][a]), "empresas": len(dados["titulares_ano"][a])} for a in anos],
        "cfem_comparable": [{"Ano": a, "v": dados["jan_jul"][a] / 100} for a in anos],
        "energy_months": anterior["energy_months"],
        "beneficiated": [{"ano": a, "venda_rs": round(dados["venda"][a], 2), "substancias": len(dados["subs_venda"][a])} for a in sorted(dados["venda"])],
        "investment": investimento(invest),
    }
    proc = processos(gpkg, args.tolerancia_processos)
    for nome, obj in [("atlas.json", packet), ("processes.json", proc)]:
        (destino / nome).write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(mun)} municípios · CFEM {dados['periodos'][0]} a {ultimo}: R$ {total_ano / 100:,.2f} · {len(subs)} substâncias em 2025 · "
          f"{proc['n']} processos ({len(base64.b64decode(proc['ringPoly'])) // 2} anéis)")
    for nome in ("atlas.json", "processes.json"):
        print(f"  {nome}: {(destino / nome).stat().st_size / 1e6:.2f} MB")


if __name__ == "__main__":
    main()
